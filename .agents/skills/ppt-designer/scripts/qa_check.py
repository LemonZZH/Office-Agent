#!/usr/bin/env python3
"""qa_check.py — PPT 交付前的程序化检查：文字溢出 / 形状越界 / 文本重叠 / 占位符残留。

用法: python3 qa_check.py deck.pptx [more.pptx ...]
退出码: 0 = 通过, 1 = 存在问题（启发式，视觉验收仍以 visual-judge 为准）
"""
import math
import re
import sys

from pptx import Presentation
from pptx.util import Emu

EMU_PER_IN = 914400
PLACEHOLDER_PAT = re.compile(r"TODO|xxx|lorem|\[insert|占位|待补充", re.IGNORECASE)


def iter_text_shapes(shapes):
    for s in shapes:
        if s.shape_type == 6:  # GROUP
            yield from iter_text_shapes(s.shapes)
        elif s.has_text_frame and s.text_frame.text.strip():
            yield s


def run_font_pt(shape, default=18.0):
    sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs
             if r.font.size is not None]
    return max(sizes) if sizes else default


def char_w(ch, pt):
    # CJK 全角按 1em，拉丁/数字按 0.55em 估宽
    return pt * (1.0 if ord(ch) > 0x2E7F else 0.55)


def est_text_height_in(shape, box_w_in):
    total_lines = 0.0
    max_pt = 0.0
    for p in shape.text_frame.paragraphs:
        text = "".join(r.text for r in p.runs) or ""
        pt = max((r.font.size.pt for r in p.runs if r.font.size), default=18.0)
        max_pt = max(max_pt, pt)
        if not text:
            total_lines += 1
            continue
        line_w = sum(char_w(c, pt) for c in text) / 72.0
        total_lines += max(1.0, math.ceil(line_w / max(box_w_in, 0.1) * 100) / 100)
    return total_lines * (max_pt or 18.0) * 1.35 / 72.0


def check(path):
    prs = Presentation(path)
    sw, sh = prs.slide_width / EMU_PER_IN, prs.slide_height / EMU_PER_IN
    issues = []
    for idx, slide in enumerate(prs.slides, 1):
        boxes = []
        for s in iter_text_shapes(slide.shapes):
            try:
                l, t = s.left / EMU_PER_IN, s.top / EMU_PER_IN
                w, h = s.width / EMU_PER_IN, s.height / EMU_PER_IN
            except TypeError:
                continue
            label = s.text_frame.text.strip().replace("\n", " ")[:24]
            if l < -0.05 or t < -0.05 or l + w > sw + 0.05 or t + h > sh + 0.05:
                issues.append(f"  p{idx} 越界: '{label}' at ({l:.2f},{t:.2f}) {w:.2f}x{h:.2f}")
            est_h = est_text_height_in(s, w - 0.1)
            if est_h > h * 1.12:
                issues.append(f"  p{idx} 溢出?: '{label}' 框高{h:.2f}\" 估需{est_h:.2f}\" "
                              f"(字号{run_font_pt(s):.0f}pt) → 缩字号/裁文案/扩框")
            if PLACEHOLDER_PAT.search(s.text_frame.text):
                issues.append(f"  p{idx} 占位符: '{label}'")
            boxes.append((l, t, w, h, label))
        for i in range(len(boxes)):
            for j in range(i + 1, len(boxes)):
                l1, t1, w1, h1, n1 = boxes[i]
                l2, t2, w2, h2, n2 = boxes[j]
                ox = min(l1 + w1, l2 + w2) - max(l1, l2)
                oy = min(t1 + h1, t2 + h2) - max(t1, t2)
                if ox > 0.05 and oy > 0.05:
                    area = ox * oy
                    smaller = min(w1 * h1, w2 * h2)
                    if smaller > 0 and area / smaller > 0.25:
                        issues.append(f"  p{idx} 重叠: '{n1}' × '{n2}' "
                                      f"({area / smaller:.0%} of smaller)")
    return issues


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    bad = False
    for path in sys.argv[1:]:
        issues = check(path)
        print(f"{'✗' if issues else '✓'} {path}" + (f" — {len(issues)} 个问题" if issues else " — 通过"))
        for it in issues:
            print(it)
        bad = bad or bool(issues)
    return 1 if bad else 0


if __name__ == "__main__":
    sys.exit(main())
